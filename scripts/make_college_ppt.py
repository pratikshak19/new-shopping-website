#!/usr/bin/env python3
"""Trendora college viva PowerPoint."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x1A, 0x1A, 0x2E)
ROSE = RGBColor(0xFF, 0x3F, 0x6C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x28, 0x2C, 0x3F)
MUTED = RGBColor(0x6B, 0x70, 0x80)
PAPER = RGBColor(0xFA, 0xF7, 0xF5)
GOLD = RGBColor(0xC9, 0xA2, 0x27)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(box, lines, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color, font)


def rect(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def tb(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h)


def title_bar(slide, kicker, title):
    rect(slide, 0, 0, W, H, PAPER)
    rect(slide, 0, 0, Inches(0.18), H, ROSE)
    k = tb(slide, Inches(0.7), Inches(0.28), Inches(12), Inches(0.35))
    add_text(k, [kicker.upper()], 12, True, ROSE)
    t = tb(slide, Inches(0.7), Inches(0.55), Inches(12), Inches(0.7))
    add_text(t, [title], 32, True, NAVY, font="Georgia")


def bullets(slide, items, top=1.4, size=20):
    box = tb(slide, Inches(0.8), Inches(top), Inches(11.6), Inches(5.6))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size, False, INK)


def cards(slide, triples, top=1.45):
    n = len(triples)
    gap = Inches(0.22)
    left = Inches(0.7)
    width = (W - Inches(1.4) - gap * (n - 1)) / n
    for i, (h, body, sub) in enumerate(triples):
        x = left + i * (width + gap)
        sh = rect(slide, x, Inches(top), width, Inches(4.6), WHITE)
        sh.shadow.inherit = False
        ht = tb(slide, x + Inches(0.2), Inches(top + 0.25), width - Inches(0.4), Inches(0.9))
        add_text(ht, [h], 18, True, ROSE)
        bd = tb(slide, x + Inches(0.2), Inches(top + 1.2), width - Inches(0.4), Inches(2.4))
        add_text(bd, [body], 15, False, INK)
        if sub:
            st = tb(slide, x + Inches(0.2), Inches(top + 3.7), width - Inches(0.4), Inches(0.6))
            add_text(st, [sub], 12, True, MUTED)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 title
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, Inches(6.85), W, Inches(0.65), ROSE)
    k = tb(s, Inches(0.9), Inches(1.6), Inches(11), Inches(0.4))
    add_text(k, ["COLLEGE PROJECT  ·  SHOPPING WEBSITE"], 14, True, ROSE)
    t = tb(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.3))
    add_text(t, ["Trendora"], 60, True, WHITE, font="Georgia")
    sub = tb(s, Inches(0.9), Inches(3.4), Inches(11), Inches(1.2))
    add_text(sub, ["Shop the trend. Live the style.", "A full mall — fashion, electronics, beauty, home and kids — with smart kits, not only clothes."], 20, False, RGBColor(0xD9, 0xD2, 0xCC))
    foot = tb(s, Inches(0.9), Inches(6.95), Inches(11), Inches(0.4))
    add_text(foot, ["Department of Computer Science  ·  Academic year 2025–26  ·  Simulated checkout"], 14, False, WHITE)

    # 2 agenda
    s = prs.slides.add_slide(blank)
    title_bar(s, "Overview", "What we will cover")
    bullets(s, [
        "Why Trendora — problem and goal",
        "Who uses the site — customer vs admin",
        "Shopping features (Myntra / Meesho style)",
        "What makes it unique — quiz, kits, festivals",
        "How it is built — React + Vite",
        "Live demo path and honest limits",
        "Conclusion and viva questions",
    ])

    # 3 problem
    s = prs.slides.add_slide(blank)
    title_bar(s, "Motivation", "The problem")
    cards(s, [
        ("Catalogue clones", "Most student shops only list products. The shopper still does not know what to buy under a budget.", "No decision help"),
        ("Too many roles", "Owner, seller, reseller logins confuse a viva. Shoppers should never add products.", "Need two doors only"),
        ("Fake payments", "Putting bank / UPI details on a college SPA does not move money and is unsafe.", "Simulate, do not charge"),
    ])

    # 4 goal
    s = prs.slides.add_slide(blank)
    title_bar(s, "Objective", "What Trendora must do")
    bullets(s, [
        "One shopping website for many departments — not a clothing boutique, not Blinkit grocery.",
        "Customer can browse, filter, bag, coupon, pay (demo) and track an order.",
        "Only a logged-in customer can place an order.",
        "Only admin can add, edit or hide products from a private /staff desk.",
        "Help the shopper decide: quiz, look of the day, festival radar, lifestyle kits.",
        "English UI, unique product photos, no repeated images.",
    ])

    # 5 inspiration
    s = prs.slides.add_slide(blank)
    title_bar(s, "Inspiration", "Habits we copied — and what we did not")
    cards(s, [
        ("Myntra", "Category shop, filters, wishlist, bag, Insider-style points, Studio content, complete-the-look.", "Fashion-first IA"),
        ("Flipkart", "Deals, coupons, COD, PIN check, order timeline Confirmed → Delivered.", "Tracking, not live map"),
        ("Meesho", "Affordable campus prices and a simple catalogue. We did not add reseller login.", "Price + clarity"),
    ])

    # 6 two accounts
    s = prs.slides.add_slide(blank)
    title_bar(s, "Users", "Only two accounts")
    cards(s, [
        ("Customer  ·  /login", "First name, last name, any email, password. Then OTP 1923. New email creates an account. Shop, bag, checkout, orders.", "demo@trendora.in / demo123"),
        ("Admin  ·  /staff", "Private desk. Add item, hide SKU, move order status, coupons, returns. Customers who type /admin are sent home.", "admin@trendora.in / admin123"),
        ("Not in this project", "Owner, seller and reseller doors were removed so the viva stays clear: shop vs catalogue control.", "Two roles only"),
    ])

    # 7 customer features
    s = prs.slides.add_slide(blank)
    title_bar(s, "Module 1", "Customer shopping website")
    bullets(s, [
        "Home, Shop by category, search, brand / rating / discount filters, sort.",
        "Product page: colour, size, PIN, reviews, questions, complete this look, add to cart / buy now.",
        "Bag, coupons (TREND10, FESTIVE20, WELCOME100, FREESHIP, INSIDER15, CAMPUS10), gift wrap.",
        "Checkout only after login → address → Razorpay-style pay sheet or COD.",
        "My orders: Confirmed → Packed → Shipped → Out for delivery → Delivered.",
        "Wishlist, compare, Insider points, help tickets, returns after delivery.",
    ])

    # 8 unique
    s = prs.slides.add_slide(blank)
    title_bar(s, "Differentiators", "Why visit Trendora — not another clone")
    bullets(s, [
        "60-second Style Quiz (/quiz) — occasion, budget, vibe, gender → complete look + add full look.",
        "Look of the day on Home — changes with the calendar, campus budget.",
        "Festival Radar (/festivals) — Indian dates (Freshers, Ganesh, Navratri, Diwali…) with days left.",
        "I already own this (/own) — extras only; owned piece is not forced into the cart.",
        "Lifestyle kits (/kits) — campus, tech, beauty, home, gift, kids mix departments under one cap.",
        "Crowd score on looks. Full mall, not clothing-only.",
    ])

    # 9 smart
    s = prs.slides.add_slide(blank)
    title_bar(s, "Style engine", "Personalized smart shopping")
    bullets(s, [
        "/style assistant: type a wish, lock a budget, pick a mood (cute, party, formal…).",
        "Complete-the-look on the product page (matching shoes, bag, extra).",
        "Why we recommend this + Eco score + Student Mode.",
        "Smart cart: “your look is almost complete — add this?”",
        "All of this stays on the same Trendora shop — we did not build a second website.",
    ])

    # 10 admin
    s = prs.slides.add_slide(blank)
    title_bar(s, "Module 2", "Private admin desk")
    bullets(s, [
        "Hidden door: footer Admin or type /staff. Public /login rejects admin.",
        "3-step Add a new item: name → price & stock → photo → Publish. It appears on Shop.",
        "Catalogue: edit or hide. Orders: advance the timeline so tracking updates for the customer.",
        "Returns, coupons, tickets, reports, store settings (announcement, free-ship minimum).",
        "Viva line: shoppers never add or delete products.",
    ])

    # 11 login checkout
    s = prs.slides.add_slide(blank)
    title_bar(s, "Flows", "Login and order")
    bullets(s, [
        "Guest may browse and fill the bag. Checkout / Buy now without login → /login?next=/checkout.",
        "After OTP 1923 the shopper returns to checkout. Bag is not lost.",
        "placeOrder refuses anyone who is not a customer.",
        "Payment: Trendora Pay sheet (UPI, cards, netbanking, wallets) or COD — demo, no charge, no bank details.",
        "Admin never needs OTP 1923. Staff password is separate.",
    ])

    # 12 tech
    s = prs.slides.add_slide(blank)
    title_bar(s, "Technology", "How it is built")
    cards(s, [
        ("Frontend", "React 18 + Vite. React Router for pages. CSS in one design language (rose on navy).", "SPA"),
        ("State", "StoreContext: users, catalogue, cart, orders in the browser (localStorage v2).", "No backend API"),
        ("Logic", "styleEngine builds looks and kits. security.js hashes passwords and checks OTP 1923.", "Client-side"),
    ])

    # 13 architecture
    s = prs.slides.add_slide(blank)
    title_bar(s, "Architecture", "One picture")
    bullets(s, [
        "Pages (Home, Shop, Quiz, Checkout, /staff Admin) talk to StoreContext.",
        "products.js is the seed catalogue with unique /images/products photos.",
        "styleEngine.js scores slots and budgets for looks / kits / festivals.",
        "RequireRole guards /admin/*. Unknown /owner and /seller routes redirect.",
        "Refresh keeps the bag because state is saved in the browser — good for a viva without a server.",
    ])

    # 14 formula
    s = prs.slides.add_slide(blank)
    title_bar(s, "Business rules", "Price and points")
    bullets(s, [
        "Grand total = subtotal − coupon + shipping + gift wrap (₹49).",
        "Free delivery when subtotal ≥ ₹999, or coupon FREESHIP.",
        "Insider: 1 point per ₹10. Elite / Icon unlock INSIDER15.",
        "Student Mode + CAMPUS10. Stock falls when an order is placed.",
        "Cancel before ship. Return / exchange only after Delivered.",
    ])

    # 15 demo
    s = prs.slides.add_slide(blank)
    title_bar(s, "Viva", "Live demo script (8 minutes)")
    bullets(s, [
        "1. Home — Why visit, Look of the day, Festival banner, Shop by category.",
        "2. Quiz — four taps, add full look, show crowd score.",
        "3. Kits — campus starter mixes bag + notes + shoes (not only clothes).",
        "4. Logout → Checkout → auto login → first name, last name, OTP 1923 → pay demo → order Confirmed.",
        "5. /staff admin123 → Add a new item → Publish → open Shop and show the new card.",
        "6. Advance order to Packed / Shipped so tracking moves.",
    ])

    # 16 testing
    s = prs.slides.add_slide(blank)
    title_bar(s, "Testing", "What we checked")
    bullets(s, [
        "Add to cart actually increases the bag count.",
        "Guest cannot place an order; customer can.",
        "Wrong OTP is rejected; 1923 works for every customer and new account.",
        "Admin cannot log in on /login; customer cannot open /admin.",
        "Women / Accessories photos match the category. No two SKUs share the same image file.",
        "Order timeline and admin status stay in sync after refresh.",
    ])

    # 17 honest
    s = prs.slides.add_slide(blank)
    title_bar(s, "Limitations", "Be honest in the viva")
    bullets(s, [
        "No real Razorpay / UPI / bank. Putting IFSC on admin would not transfer money.",
        "OTP is simulated (always 1923). No SMS gateway.",
        "Passwords are hashed in the browser — still not a bank-grade server.",
        "localStorage is a teaching stand-in for a REST API.",
        "Nothing ships. This is an academic shopping-website project.",
        "To go live: backend + Razorpay merchant KYC + host on Vercel after GitHub push.",
    ])

    # 18 future
    s = prs.slides.add_slide(blank)
    title_bar(s, "Future scope", "If this were a product")
    bullets(s, [
        "Node / Spring API and a real database.",
        "Razorpay Orders API + webhook → Confirmed only after paid.",
        "Cloud images and admin photo upload.",
        "Email / SMS OTP from a provider.",
        "PWA install and better mobile nav.",
    ])

    # 19 conclusion
    s = prs.slides.add_slide(blank)
    title_bar(s, "Conclusion", "One sentence for the examiner")
    box = tb(s, Inches(0.9), Inches(1.8), Inches(11.4), Inches(3.5))
    add_text(box, [
        "Trendora is a full shopping website, not a clothing boutique and not a quick-commerce app.",
        "",
        "It copies the useful habits of Myntra and Flipkart, then adds decision tools — quiz, kits, festivals — while keeping a private admin desk and a simulated, safe checkout for college.",
    ], 22, False, INK)

    # 20 thank you
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, Inches(6.85), W, Inches(0.65), ROSE)
    t = tb(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.2))
    add_text(t, ["Thank you"], 54, True, WHITE, font="Georgia")
    q = tb(s, Inches(0.9), Inches(3.5), Inches(11), Inches(1.8))
    add_text(q, [
        "Questions?",
        "Demo: /login  ·  OTP 1923   |   Admin: /staff  ·  admin123",
        "Quiz  ·  Kits  ·  Festivals  ·  /documents",
    ], 18, False, RGBColor(0xD9, 0xD2, 0xCC))
    f = tb(s, Inches(0.9), Inches(6.95), Inches(11), Inches(0.4))
    add_text(f, ["Trendora  ·  Shop the trend. Live the style."], 14, False, WHITE)

    out = [
        "/home/user/new-shopping-website/COLLEGE_DOCUMENTS/Trendora_College_Presentation.pptx",
        "/home/user/new-shopping-website/DOWNLOAD_THESE/Trendora_College_Presentation.pptx",
        "/home/user/new-shopping-website/public/docs/Trendora_College_Presentation.pptx",
    ]
    for path in out:
        prs.save(path)
        print("saved", path)


if __name__ == "__main__":
    build()
